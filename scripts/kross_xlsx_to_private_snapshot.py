#!/usr/bin/env python3
"""Build a dated, owner-only OpenClaw snapshot from a Kross XLSX export.

The output is a provisional operational projection, never a live availability
feed. The destination must already exist so its private filesystem permissions
are preserved while writing.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import stat
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAX_RESERVATIONS = 500
MAX_BYTES = 524_288
SOURCE_ID_RE = re.compile(r"^(\d{1,12})/(\d{4})$")
ALLOWED_STATUSES = frozenset(
    {
        "Pendiente de Pago",
        "Reserva Nueva (Revisar)",
        "Confirmada (Lista para Check In)",
        "Paga al Check-In",
        "Check-in",
        "Extras pendientes",
    }
)
REQUIRED_HEADERS = frozenset(
    {
        "Number",
        "Name",
        "Rooms",
        "Guests",
        "Arrival",
        "Departure",
        "Channel",
        "Status",
        "Total",
    }
)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--base-snapshot", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--cutoff-date", required=True, type=date.fromisoformat)
    parser.add_argument("--source-as-of", required=True)
    parser.add_argument("--expected-sha256", required=True)
    return parser.parse_args(argv)


def strict_source_as_of(value: str) -> str:
    if not re.fullmatch(
        r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?(?:Z|[+-]\d{2}:\d{2})",
        value,
    ):
        raise ValueError("source-as-of must be an ISO-8601 timestamp with timezone")
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("source-as-of must include timezone")
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: Any, *, max_length: int = 500) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    if not text:
        return None
    if len(text) > max_length:
        raise ValueError("cell text exceeds safe length")
    return text


def as_date(value: Any, *, row_number: int, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    raise ValueError(f"row {row_number}: invalid {field}")


def optional_count(value: Any, *, row_number: int) -> int | None:
    if value is None or clean_text(value) is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"row {row_number}: invalid Guests")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row_number}: invalid Guests") from error
    if not math.isfinite(number) or number < 0 or not number.is_integer():
        raise ValueError(f"row {row_number}: invalid Guests")
    return int(number)


def optional_amount(value: Any, *, row_number: int) -> float | None:
    if value is None or clean_text(value) is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"row {row_number}: invalid Total")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row_number}: invalid Total") from error
    if not math.isfinite(number):
        raise ValueError(f"row {row_number}: invalid Total")
    return round(number, 2)


def room_projection(value: Any) -> tuple[list[dict[str, str]], str]:
    text = clean_text(value, max_length=200)
    if text is None or text.upper() == "N/A":
        return [], "UNKNOWN"
    rooms = []
    for item in text.split(","):
        room = clean_text(item, max_length=80)
        if room:
            rooms.append({"room": room})
    return rooms, "PRESENT" if rooms else "UNKNOWN"


def load_base(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    required = {"orgId", "propertyId", "operations", "exportedAt", "sourceIsLive"}
    if not required.issubset(payload):
        raise ValueError("base snapshot is missing required fields")
    if payload["sourceIsLive"] is not False:
        raise ValueError("base snapshot must be non-live")
    operations_at = payload.get("operations", {}).get("capturedAt")
    strict_source_as_of(operations_at)
    return payload


def build_snapshot(
    input_path: Path,
    base_snapshot: dict[str, Any],
    *,
    cutoff_date: date,
    source_as_of: str,
    source_hash: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as error:
        raise ValueError("workbook is empty") from error
    headers = [clean_text(value, max_length=100) or "" for value in raw_headers]
    indexes = {header: index for index, header in enumerate(headers)}
    missing = sorted(REQUIRED_HEADERS - indexes.keys())
    if missing:
        raise ValueError(f"workbook is missing required headers: {', '.join(missing)}")

    reservations: list[dict[str, Any]] = []
    normalized_ids: set[str] = set()
    source_ids: set[str] = set()
    excluded_statuses: Counter[str] = Counter()
    total_rows = 0

    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None for value in row):
            continue
        total_rows += 1
        departure = as_date(row[indexes["Departure"]], row_number=row_number, field="Departure")
        status_value = clean_text(row[indexes["Status"]], max_length=100)
        if departure < cutoff_date or status_value not in ALLOWED_STATUSES:
            excluded_statuses[status_value or "UNKNOWN"] += 1
            continue

        arrival = as_date(row[indexes["Arrival"]], row_number=row_number, field="Arrival")
        if departure < arrival:
            raise ValueError(f"row {row_number}: departure precedes arrival")
        source_id = clean_text(row[indexes["Number"]], max_length=40)
        match = SOURCE_ID_RE.fullmatch(source_id or "")
        if not match:
            raise ValueError(f"row {row_number}: invalid Kross Number")
        normalized_id = match.group(1)
        if source_id in source_ids or normalized_id in normalized_ids:
            raise ValueError(f"row {row_number}: duplicate reservation identifier")
        source_ids.add(source_id)
        normalized_ids.add(normalized_id)

        guest_name = clean_text(row[indexes["Name"]], max_length=200)
        if guest_name is None:
            raise ValueError(f"row {row_number}: missing Name")
        guests = optional_count(row[indexes["Guests"]], row_number=row_number)
        channel = clean_text(row[indexes["Channel"]], max_length=100)
        total = optional_amount(row[indexes["Total"]], row_number=row_number)
        rooms, rooms_status = room_projection(row[indexes["Rooms"]])

        reservations.append(
            {
                "reservationId": normalized_id,
                "guestName": guest_name,
                "phone": None,
                "email": None,
                "plannedArrival": arrival.isoformat(),
                "plannedDeparture": departure.isoformat(),
                "statusPms": status_value,
                "guestsCount": guests,
                "adults": None,
                "children": None,
                "channel": channel,
                "breakfastIncluded": None,
                "totalAmount": total,
                "currency": None,
                "amountPaid": None,
                "paymentStatus": None,
                "recordedPayments": [],
                "snapshotAsOf": source_as_of,
                "details": {
                    "schemaVersion": "kross-xlsx-operational-v1",
                    "rooms": rooms,
                    "source": f"Kross dated XLSX; source reservation {source_id}",
                    "sourceReservationId": source_id,
                    "sourceAsOf": source_as_of,
                    "coverage": {
                        "rooms": rooms_status,
                        "guests": "PRESENT" if guests is not None else "UNKNOWN",
                        "channel": "PRESENT" if channel is not None else "UNKNOWN",
                        "total": "PRESENT" if total is not None else "UNKNOWN",
                        "payments": "NOT_IN_EXPORT",
                        "breakfast": "NOT_IN_EXPORT",
                    },
                    "phoneStatus": "NOT_IN_EXPORT",
                    "emailStatus": "NOT_IN_EXPORT",
                    "currencyStatus": "NOT_IN_EXPORT",
                    "breakfastStatus": "NOT_IN_EXPORT",
                    "financialStatus": "TOTAL_ONLY_PAYMENT_UNKNOWN",
                    "actualStayEvidence": {"status": "UNKNOWN"},
                    "internalNoteExcluded": True,
                    "noteSummaryRequiresReview": False,
                },
            }
        )

    workbook.close()
    reservations.sort(key=lambda item: (item["plannedArrival"], item["reservationId"]))
    if not reservations:
        raise ValueError("projection contains no reservations")
    if len(reservations) > MAX_RESERVATIONS:
        raise ValueError("projection exceeds the reader reservation limit")

    snapshot = {
        "orgId": base_snapshot["orgId"],
        "propertyId": base_snapshot["propertyId"],
        "source": "Kross dated XLSX provisional private operational projection",
        "exportedAt": source_as_of,
        "sourceIsLive": False,
        "reservations": reservations,
        "operations": base_snapshot["operations"],
        "provenance": {
            "schemaVersion": "kross-xlsx-private-snapshot-v1",
            "sourceType": "local_xlsx_export",
            "sourceFileName": input_path.name,
            "sourceSha256": source_hash,
            "sourceAsOf": source_as_of,
            "cutoffDate": cutoff_date.isoformat(),
            "sourceRowCount": total_rows,
            "selectedReservationCount": len(reservations),
            "reservationIdStrategy": "numeric prefix for current reader; exact source ID retained privately",
            "provenanceStatus": "attributed_not_cryptographically_verified",
        },
    }
    report = {
        "status": "BUILT",
        "sourceRows": total_rows,
        "selectedReservations": len(reservations),
        "excludedByStatus": dict(sorted(excluded_statuses.items())),
        "sourceAsOf": source_as_of,
        "sourceSha256": source_hash,
        "sourceIsLive": False,
    }
    return snapshot, report


def encode_snapshot(snapshot: dict[str, Any]) -> bytes:
    data = (json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if len(data) > MAX_BYTES:
        raise ValueError("snapshot exceeds the reader byte limit")
    return data


def write_existing_private_output(path: Path, data: bytes) -> None:
    output_stat = path.lstat()
    if not stat.S_ISREG(output_stat.st_mode) or path.is_symlink():
        raise ValueError("output must be an existing regular file")
    if output_stat.st_size != 0:
        raise ValueError("output must be an empty pre-created private file")
    with path.open("r+b") as handle:
        handle.write(data)
        handle.truncate()
        handle.flush()
        os.fsync(handle.fileno())


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    source_as_of = strict_source_as_of(args.source_as_of)
    actual_hash = sha256_file(args.input)
    if actual_hash.lower() != args.expected_sha256.lower():
        raise ValueError("input SHA-256 does not match the approved artifact")
    base_snapshot = load_base(args.base_snapshot)
    snapshot, report = build_snapshot(
        args.input,
        base_snapshot,
        cutoff_date=args.cutoff_date,
        source_as_of=source_as_of,
        source_hash=actual_hash,
    )
    data = encode_snapshot(snapshot)
    write_existing_private_output(args.output, data)
    report["bytes"] = len(data)
    report["outputSha256"] = hashlib.sha256(data).hexdigest()
    print(json.dumps(report, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
